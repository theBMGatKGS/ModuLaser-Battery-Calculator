# 02.00.000 checks: new layout, Print/Export (PDF/DOCX/HTML/Excel x 3 detail levels), Help audit, screenshots.
import asyncio, json, os, re, zipfile, subprocess, shutil
from playwright.async_api import async_playwright
from helpers import *
import os as _os, tempfile as _tf
MLP_OUT=_os.environ.get('MLP_TEST_OUT') or _tf.mkdtemp(prefix='mlp-test-')   # all generated output goes here (QA r3)
def _od(*p):
    d=_os.path.join(MLP_OUT,*p); _os.makedirs(d,exist_ok=True); return d.rstrip('/')+'/'
SH=_od('calcshots2'); OUT=_od('calcout'); os.makedirs(SH,exist_ok=True); os.makedirs(OUT,exist_ok=True)
res=[]
def ok(name,cond,info=''):
    res.append((name,bool(cond),info)); print(('PASS' if cond else 'FAIL'),name,str(info)[:300])
LEVELS=['summary','standard','full']
STEP_TITLES={'full':['1. Module & supply load','2. Standby & alarm capacity','3. Safety margin & temperature','4. Battery & supply sizing'],
             'standard':['1. Standby & alarm capacity','2. Safety margin & temperature','3. Battery & supply sizing']}
async def dl(pg, sel):
    async with pg.expect_download() as d: await pg.click(sel)
    dd=await d.value; p=OUT+dd.suggested_filename; await dd.save_as(p); return p
async def main():
    async with async_playwright() as p:
        b=await p.chromium.launch(); ctx=await b.new_context(viewport={'width':1500,'height':1000}, accept_downloads=True)
        pg=await ctx.new_page(); pg.set_default_timeout(120000)
        errs=[]; pg.on('pageerror',lambda e: errs.append(str(e)))
        await pg.goto(APP); await pg.evaluate('localStorage.clear()'); await pg.reload(); await pg.wait_for_timeout(500)
        ok('no CDN / external script', await pg.evaluate("[...document.scripts].every(s=>!s.src)"))
        await import_files(pg,['BWDG1.pl','Solano Prison Vacaville CA.pl','Factory Motor Parts ASRS Main Room (OAP).pl']); await fill_addresses(pg)
        await pg.evaluate("v2Refresh()")
        n=await pg.evaluate("document.querySelectorAll('.location').length")
        ok('sidebar lists every sheet', await pg.evaluate("document.querySelectorAll('.sb-item').length")==n, n)
        ok('sidebar grouped by Area/Group', await pg.evaluate("document.querySelectorAll('.sb-gh').length")>=3)
        ok('only current sheet visible', await pg.evaluate("[...document.querySelectorAll('.location')].filter(l=>getComputedStyle(l).display!=='none').length")==1)
        await pg.locator('.sb-item').nth(2).click(); await pg.wait_for_timeout(150)
        ok('sidebar click switches sheet', await pg.evaluate("document.querySelector('.location-current').dataset.loc")==await pg.evaluate("document.querySelectorAll('.sb-item')[2].dataset.loc"))
        # results panel matches on-screen calc
        rp=await pg.evaluate("({big:document.querySelector('.rp-big').textContent, req:document.querySelector('.location-current .required-ah').textContent, st:document.getElementById('rp-status').textContent})")
        ok('results panel Ah matches sheet', abs(float(rp['big'].split()[0])-float(rp['req'].split()[0]))<0.06, rp)
        # status dots incl. error + overridden
        await pg.evaluate("(()=>{const l=document.querySelector('.location-current');l.querySelector('.standby-hours').value='48';l.querySelector('.design-basis-select').value='Custom';applyDesignBasis(l.dataset.loc);l.querySelector('.standby-hours').value='48';recalcLocation(l.dataset.loc);})()")
        await pg.evaluate("(()=>{const l=document.querySelectorAll('.location')[0];l.querySelector('[data-role=module] .module-address').value='';recalcLocation(l.dataset.loc);})()")
        await pg.evaluate("v2Refresh()"); await pg.wait_for_timeout(200)
        dots=await pg.evaluate("[...document.querySelectorAll('.sb-item .dot')].map(d=>d.className.replace('dot st-',''))")
        ok('status dots: complete/needs/error/overridden all shown', {'error','overridden'}<=set(dots) and ({'complete','needs'}&set(dots)), sorted(set(dots)))
        ok('override badge on sheet', await pg.evaluate("!document.querySelector('.location-current .v2-or-badge').hidden && /Standby 48/.test(document.querySelector('.location-current .v2-or-badge').textContent)"))
        await pg.screenshot(path=SH+'01-main-screen.png')
        await pg.locator('#sheet-sidebar').screenshot(path=SH+'03-sidebar-statuses.png')
        # 16-char counter
        await pg.evaluate("(()=>{const i=document.querySelector('.location-current .zone-name');i.value='ABCDEFGHIJKLMNOPQ';i.dispatchEvent(new Event('input',{bubbles:true}));})()")
        ok('n/16 counter red over 16', await pg.evaluate("(()=>{const c=document.querySelector('.location-current .v2-cc');return c.textContent==='17/16'&&c.classList.contains('over')})()"))
        await pg.evaluate("(()=>{const i=document.querySelector('.location-current .zone-name');i.value='Zone A';i.dispatchEvent(new Event('input',{bubbles:true}));})()")
        # settings bar collapse
        await pg.click('#proj-settings-toggle'); await pg.wait_for_timeout(100)
        ok('Project Settings collapses to one-line summary', await pg.evaluate("document.getElementById('proj-settings').classList.contains('collapsed') && /NDA Client Expansion/.test(document.getElementById('proj-settings-summary').textContent)"))
        await pg.screenshot(path=SH+'01b-settings-collapsed.png', clip={'x':0,'y':0,'width':1500,'height':260})
        await pg.click('#proj-settings-toggle'); await pg.wait_for_timeout(100)
        await pg.screenshot(path=SH+'02-settings-expanded.png', clip={'x':0,'y':0,'width':1500,'height':330})
        # narrow
        await pg.set_viewport_size({'width':760,'height':1000}); await pg.wait_for_timeout(300)
        ok('narrow: side panels collapse', await pg.evaluate("getComputedStyle(document.getElementById('app-shell')).gridTemplateColumns.split(' ').length<=2"))
        await pg.screenshot(path=SH+'01c-narrow.png'); await pg.set_viewport_size({'width':1500,'height':1000}); await pg.wait_for_timeout(200)
        # restore a clean state for exports (undo forced error, keep override)
        await pg.evaluate("(()=>{const l=document.querySelectorAll('.location')[0];const s=l.querySelector('[data-role=module] .module-address');s.value=s.dataset.desiredValue||'1';recalcLocation(l.dataset.loc);v2Refresh();})()")
        # ---- Print/Export dialog ----
        await pg.click('#print-btn'); await pg.wait_for_timeout(600)
        ok('Print opens Print/Export dialog', await pg.evaluate("document.getElementById('pe-overlay').classList.contains('open')"))
        models=await pg.evaluate("v2AllModels('all').map(m=>({tag:v2PsTag(m),req:m.requiredAh.toFixed(2),st:v2StatusWord(m.status),n:m.modules.length}))")
        for lv in LEVELS:
            await pg.check(f'#pe-level-{lv}'); await pg.wait_for_timeout(700)
            fr=pg.frame_locator('#pe-preview')
            await pg.screenshot(path=SH+f'04-print-dialog-{lv}.png')
            html=await pg.evaluate(f"v2BuildPrintHtml(Object.assign(v2PeOpts(),{{level:'{lv}'}}))")
            pdfp=OUT+f'report-{lv}.pdf'; pages,npg,over=await render_pdf(b, html, pdfp)
            ok(f'[{lv}] PDF renders, pages == layout pages', pages==npg, f'{pages} pdf pages')
            ok(f'[{lv}] every page (PSU + appendix) fits its letter page', not over, over)
            psu=html.count('class="page psu"'); ok(f'[{lv}] one PSU page per sheet', psu==len(models), psu)
            # values match screen
            miss=[m['tag'] for m in models if f"Required Ah: {m['req']} Ah" not in html]
            ok(f'[{lv}] printed Required Ah matches on-screen results', not miss, miss)
            if lv in STEP_TITLES:
                t=[s for s in STEP_TITLES[lv] if html.count(s.replace('&','&amp;'))<len(models)]
                ok(f'[{lv}] every calculation step present for every PSU', not t, t)
            if lv=='full':
                ok('[full] appendix page per PSU', html.count('class="page appendix"')>=len(models))
                ok('[full] per-module currents + formula/substituted/result/reference', 'FAN_SPEED_MA[' in html and 'Standby Ah = I_total × T_standby' in html and 'NFPA 72 (2022 Edition), Section 10.6.7.2' in html)
            # HTML
            hp=await dl(pg,'#pe-html'); hh=open(hp).read()
            ok(f'[{lv}] HTML export self-contained', '<script src' not in hh and '<link' not in hh and 'class="page psu"' in hh, os.path.basename(hp)); shutil.copy(hp,OUT+f'report-{lv}.html')
            if lv=='full': ok('[full] HTML uses expandable sections', '<details' in hh)
            # DOCX
            dp=await dl(pg,'#pe-docx'); shutil.move(dp,OUT+f'report-{lv}.docx'); dp=OUT+f'report-{lv}.docx'
            z=zipfile.ZipFile(dp); doc=z.read('word/document.xml').decode()
            ok(f'[{lv}] DOCX valid package with PSU pages + images', z.testzip() is None and doc.count('Power Supply ')>=len(models) and any(n.startswith('word/media/') for n in z.namelist()))
            if lv in STEP_TITLES: ok(f'[{lv}] DOCX has every step', all(doc.count(s.replace('&','&amp;'))>=len(models) for s in STEP_TITLES[lv]))
            r=subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',OUT+'docx/',dp],capture_output=True,text=True,timeout=180)
            dpdf=OUT+'docx/'+os.path.basename(dp)[:-5]+'.pdf'; dinfo=subprocess.run(['pdfinfo',dpdf],capture_output=True,text=True).stdout if os.path.exists(dpdf) else ''
            ok(f'[{lv}] DOCX opens in LibreOffice (letter pages)', '612 x 792' in dinfo, [l for l in dinfo.splitlines() if l.startswith(('Pages','Page size'))])
            # Excel
            xp=await dl(pg,'#pe-xlsx'); shutil.move(xp,OUT+f'report-{lv}.xlsx'); xp=OUT+f'report-{lv}.xlsx'
            import openpyxl; wb=openpyxl.load_workbook(xp)
            ok(f'[{lv}] Excel has CAD Schedule + Detail', wb.sheetnames[:2]==['CAD Schedule','Detail'], wb.sheetnames)
            ok(f'[{lv}] Excel has no merged cells', all(not ws.merged_cells.ranges for ws in wb.worksheets))
            cad=wb['CAD Schedule']; fills=[c.coordinate for row in cad.iter_rows() for c in row if c.fill and c.fill.fill_type not in (None,'none')]
            ok(f'[{lv}] CAD Schedule: no fills, fixed widths, footer per PSU', not fills and cad.column_dimensions['A'].width==24 and sum(1 for r in cad.iter_rows(values_only=True) if str(r[0]).startswith('Required Ah'))==len(models))
            det=wb['Detail']; lvls=[det.row_dimensions[i].outlineLevel for i in range(1,det.max_row+1)]
            ok(f'[{lv}] Detail has outline row groups', max(lvls)>=(1 if lv=='summary' else 2), f'max level {max(lvls)}, rows {det.max_row}')
            if lv=='full': ok('[full] Detail shows every step', all(any(r[0]==s for r in det.iter_rows(values_only=True)) for s in STEP_TITLES['full']))
            open(OUT+f'cad-{lv}.json','w').write(json.dumps([list(r) for r in cad.iter_rows(values_only=True)]))
        ok('CAD Schedule identical at every level', len({open(OUT+f'cad-{lv}.json').read() for lv in LEVELS})==1)
        await pg.check('#pe-level-standard'); await pg.wait_for_timeout(600); await pg.screenshot(path=SH+'04-print-dialog.png')
        await pg.click('#pe-close')
        await pg.click('#menu-file-btn'); await pg.click('#print-menu-btn'); await pg.wait_for_timeout(300)
        ok('File > Print / Export opens dialog', await pg.evaluate("document.getElementById('pe-overlay').classList.contains('open')")); await pg.click('#pe-close')
        await pg.keyboard.press('Control+p'); await pg.wait_for_timeout(300)
        ok('Ctrl+P opens dialog', await pg.evaluate("document.getElementById('pe-overlay').classList.contains('open')")); await pg.click('#pe-close')
        # ---- Help audit ----
        audit=await pg.evaluate('''()=>{
          const help=document.getElementById('help-overlay').textContent.replace(/\\s+/g,' ');
          const labels=[...document.querySelectorAll('.menubar-menu [role=menuitem], .menubar-menu [role=menuitemcheckbox], .quick-row button, #sb-add-btn, .pe-side button, .pe-side h3, .pe-side label > span')]
            .map(e=>(e.childNodes[0]&&e.childNodes[0].nodeType===3?e.childNodes[0].textContent:e.textContent).trim()).filter(Boolean);
          const norm=s=>s.replace(/\\s+/g,' ').trim();
          const missing=[...new Set(labels)].filter(l=>!help.includes(norm(l)));
          // reverse: every menu/button the Help names in the menu-bar step must exist
          const step=document.getElementById('qs-step-menubar').parentNode.textContent;
          const named=['New…','Open / Import Calc…','Import .pl…','Linked .pl files…','Save As…','Export to Excel','Print / Export…','Print BOM','+ Add Location','Clear…','Confirm defaults','Clear saved data','Cluster Report','Project Settings bar','Sheet sidebar','Results panel','Help','How Do I…','Revision log','Import .pl','Print'];
          const all=[...document.querySelectorAll('button,[role=menuitemcheckbox]')].flatMap(b=>[(b.childNodes[0]&&b.childNodes[0].nodeType===3?b.childNodes[0].textContent:'').trim(), b.textContent.trim()]);
          const ghost=named.filter(n=>step.includes(n)&&!all.includes(n));
          const qs=document.getElementById('help-howtouse').textContent;
          return {missing, ghost, toolbar:(qs.match(/toolbar/gi)||[]).length, labels:labels.length};}''')
        ok('Help audit: every menu item/button/dialog option is documented', not audit['missing'], audit['missing'])
        ok('Help audit: every command Help names exists', not audit['ghost'], audit['ghost'])
        ok('Help audit: no stale toolbar references in How to Use', audit['toolbar']==0, audit['toolbar'])
        hd=await pg.evaluate("(buildHowDoIIndex(), ['print-export','detail-levels','cad-paste','sidebar','results-panel','project-settings','layout'].filter(id=>!howdoiIndex.some(i=>i.id===id)))")
        ok('How Do I… has the new topics', not hd, hd)
        tgt=await pg.evaluate("HOWDOI_COMMON.filter(c=>!document.getElementById(c.target)).map(c=>c.id)"); ok('How Do I… targets exist', not tgt, tgt)
        await pg.click('#menu-help-btn'); await pg.click('#help-btn'); await pg.wait_for_timeout(300)
        await pg.evaluate("document.getElementById('qs-step-layout').scrollIntoView()"); await pg.wait_for_timeout(150); await pg.screenshot(path=SH+'09-help-quickstart.png')
        await pg.evaluate("document.getElementById('qs-step-detail-levels').scrollIntoView()"); await pg.wait_for_timeout(150); await pg.screenshot(path=SH+'09b-help-detail-levels.png')
        await pg.click('#help-close-btn')
        await pg.click('#menu-help-btn'); await pg.click('#revision-log-menu-btn'); await pg.wait_for_timeout(300)
        await pg.screenshot(path=SH+'10-help-revision-log.png')
        rl=await pg.evaluate("document.querySelector('#about-amendments-body tr td').textContent")
        is_tb = (await pg.evaluate("runningRevisionFull()")).startswith('00.')   # QA F4: the Toolbox's own log is 00.x
        if is_tb: ok('revision log top entry 00.02.000 (Toolbox)', rl=='00.02.000', rl)
        else: ok('revision log top entry 02.01.000, older entries keep 01.x', rl=='02.01.000' and await pg.evaluate("[...document.querySelectorAll('#about-amendments-body tr td:first-child')].some(t=>t.textContent==='01.05.009')") and not await pg.evaluate("[...document.querySelectorAll('#about-amendments-body tr td:first-child')].some(t=>t.textContent.startsWith('01.06'))"))
        await pg.click('#help-close-btn')
        # wizard step 4 screenshot
        await pg.click('#import-pl-btn'); await pick_step1(pg); await pg.click('#plw-next'); await pg.set_input_files('#plw-file',[T+'Deck MDA-A-D2.pl',T+'BWDG1.pl']); await pg.wait_for_timeout(600)
        await pg.click('#plw-next'); await pg.click('#plw-suggest-all'); await pg.click('#plw-next'); await pg.wait_for_timeout(300)
        ok('wizard step 4 tray + Pack into fewest present', await pg.evaluate("!!document.querySelector('.plw-tray') && [...document.querySelectorAll('[data-pack]')].some(b=>/Pack into fewest/.test(b.textContent))"))
        await pg.locator('.plw-tray').first.scroll_into_view_if_needed(); await pg.screenshot(path=SH+'05-wizard-step4.png')
        await pg.click('#plw-cancel')
        ok('no page errors (v2)', not errs, errs[:3])
        await b.close()
    import glob
    def pg_png(pdf, page, dest, r=80):
        subprocess.run(['pdftoppm','-r',str(r),'-png','-f',str(page),'-l',str(page),'-singlefile',pdf,dest[:-4]],check=True)
    pg_png(OUT+'report-standard.pdf',1,SH+'06-print-summary-page.png')
    pg_png(OUT+'report-standard.pdf',2,SH+'07-print-psu-page.png')
    for lv in LEVELS: pg_png(OUT+f'report-{lv}.pdf',2,SH+f'07-level-{lv}-psu-page.png')
    pg_png(OUT+'report-standard.pdf',3,SH+'07-level-standard-appendix.png')
    pg_png(OUT+'report-full.pdf',3,SH+'07-level-full-appendix.png')
    pg_png(OUT+'docx/report-full.pdf',2,SH+'11-docx-full-psu-page.png',70)
    pg_png(OUT+'docx/report-full.pdf',3,SH+'11b-docx-full-appendix.png',70)
    subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',OUT+'xl/',OUT+'report-full.xlsx'],capture_output=True,timeout=180)
    pg_png(OUT+'xl/report-full.pdf',1,SH+'08-excel-cad-schedule.png',90)
    print('SUMMARY', sum(1 for r in res if r[1]),'/',len(res),'passed')
    return res
if __name__=='__main__': asyncio.run(main())
